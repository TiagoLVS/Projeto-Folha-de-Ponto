import hashlib
import re
import unicodedata
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from backend.database.calendario import importar_datas
from backend.models.calendario import DataEspecialPayload, TipoData

MAX_LINHAS = 5000


def interpretar_data(value):
    if isinstance(value, datetime):
        if value.time() != time():
            raise ValueError('A célula deve conter somente uma data, sem horário.')
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r'\d{2}/\d{2}/\d{4}', value.strip()):
        try:
            return datetime.strptime(value.strip(), '%d/%m/%Y').date()
        except ValueError:
            pass
    raise ValueError('Data inválida: use uma célula de data do Excel ou DD/MM/AAAA.')


def sugerir_tipo(text):
    simple = ''.join(c for c in unicodedata.normalize('NFD', text.lower()) if not unicodedata.combining(c))
    for pattern, tipo in ((r'^ponto facultativo\b', TipoData.PONTO_FACULTATIVO), (r'^recesso\b', TipoData.RECESSO), (r'^feriado\b', TipoData.FERIADO)):
        if re.search(pattern, simple):
            return tipo
    return TipoData.OUTRO


def ler_calendario(caminho):
    try:
        with ZipFile(caminho) as archive:
            if sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError('Planilha expandida muito grande (limite de 64 MB).')
        workbook = load_workbook(caminho, read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError, InvalidFileException, ParseError) as error:
        raise ValueError('Arquivo Excel inválido ou corrompido.') from error
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError('A planilha não possui uma aba ativa.')
        rows = sheet.iter_rows()
        header = next(rows, ())
        names = [' '.join(str(cell.value or '').split()).lower() for cell in header]
        for name in ('data', 'especificação'):
            if names.count(name) != 1:
                raise ValueError('A primeira linha deve conter uma única coluna Data e uma única coluna Especificação.')
        data_idx, text_idx = names.index('data'), names.index('especificação')
        records = []
        for index, row in enumerate(rows, 2):
            if index > MAX_LINHAS + 1:
                raise ValueError(f'A planilha deve conter no máximo {MAX_LINHAS} linhas de dados.')
            values = [cell.value for cell in row]
            if all(value is None or value == '' for value in values):
                continue
            try:
                raw_date, raw_text = row[data_idx], row[text_idx]
                if raw_date.data_type == 'f' or raw_text.data_type == 'f':
                    raise ValueError('Fórmulas não são aceitas. Cole os valores da data e especificação.')
                day = interpretar_data(raw_date.value)
                if not isinstance(raw_text.value, str) or not raw_text.value.strip():
                    raise ValueError('Especificação não preenchida ou não textual.')
                event = DataEspecialPayload(data=day, especificacao=raw_text.value, tipo=sugerir_tipo(raw_text.value.strip()))
                records.append(dict(linha=index, evento=event))
            except (ValueError, ValidationError) as error:
                message = error.errors()[0]['msg'] if isinstance(error, ValidationError) else str(error)
                records.append(dict(linha=index, erro=message))
        return records
    except (ParseError, BadZipFile, KeyError) as error:
        raise ValueError('Arquivo Excel inválido ou corrompido.') from error
    finally:
        workbook.close()


def importar_calendario(caminho, nome_arquivo):
    records = ler_calendario(caminho)
    digest = hashlib.sha256(Path(caminho).read_bytes()).hexdigest()
    return importar_datas(records, nome_arquivo, digest)
